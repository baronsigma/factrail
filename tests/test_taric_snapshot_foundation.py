from __future__ import annotations

import json
import sqlite3
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook

from factrail.trade.official_taric_store import (
    FULL_REQUIRED_TABLES,
    OfficialTaricStore,
    TaricSnapshotError,
)


def _official_shaped_workbook(path: Path, *, corrupt: bool = False) -> Path:
    """Small synthetic TEST DATA; rows are shaped after the Commission guide."""
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)

    sheet("Goods_Nomenclature", ["Goods code", "Hierarchy level", "Is leaf", "Validity start", "Validity end", "Description"], [
        ["8508", "4", "0", "2020-01-01", "", "Test heading"],
        ["850870", "6", "0", "2020-01-01", "", "Test subheading"],
        ["85087000", "8", "0", "2020-01-01", "", "Test CN"],
        ["8508700010", "10", "1", "2020-01-01", "", "Synthetic test leaf"],
        ["8508700011", "10", "1", "2020-01-01", "", "Synthetic test sibling"],
    ])
    sheet("Declarable_Codes", ["Nomenclature code", "Is leaf", "Validity start", "Validity end"], [
        ["850800000000", "0", "2020-01-01", ""], ["850870000000", "0", "2020-01-01", ""],
        ["850870000080", "0", "2020-01-01", ""], ["850870001080", "1", "2020-01-01", ""],
        ["850870001180", "1", "2020-01-01", ""],
    ])
    sheet("Measures", ["Goods code", "Additional code", "Quota order number", "Validity start", "Validity end", "Geographical area", "Measure type", "Duty expression", "Regulation"], [
        ["8508", "", "", "2020-01-01", "2030-12-31", "G1", "103", "4.5 %", "R-ACT-1"],
        ["8508700010", "", "", "2020-01-01", "2024-12-31", "G1", "103", "9 %", "OLD"],
        ["8508700010", "X123", "", "2020-01-01", "2030-12-31", "G1", "551", "60 EUR/tonne", "R-ACT-2"],
    ])
    sheet("Geographical_Areas", ["Area ID", "Description"], [["G1", "Synthetic test geography group"]])
    sheet("Geographical_Area_Membership", ["Parent area", "Country code", "Validity start", "Validity end"], [["G1", "CN", "2020-01-01", "2030-12-31"], ["G1", "US", "2020-01-01", "2030-12-31"]])
    sheet("Measure_Exclusions", ["Goods code", "Measure type", "Geographical area", "Excluded area", "Validity start", "Validity end"], [["8508", "103", "G1", "CN", "2020-01-01", "2030-12-31"]])
    sheet("Measure_Conditions", ["Goods code", "Measure type", "Geographical area", "Condition code", "Certificate code", "Duty amount", "Measurement unit", "Action"], [["8508700010", "551", "G1", "Y", "C001", "", "", "Submit certificate"]])
    sheet("Measure_Footnotes", ["Goods code", "Measure type", "Geographical area", "Footnote code"], [["8508700010", "551", "G1", "TN001"]])
    sheet("Legal_Bases", ["Regulation", "Description"], [["R-ACT-1", "Synthetic legal reference"], ["R-ACT-2", "Synthetic additional measure reference"]])
    sheet("Additional_Codes", ["Additional code", "Description"], [["X123", "Synthetic additional code"]])
    sheet("Certificates", ["Certificate code", "Description"], [["C001", "Synthetic certificate requirement"]])
    sheet("Footnotes", ["Footnote code", "Description"], [["TN001", "Synthetic test footnote"]])
    sheet("Measurement_Units", ["Measurement unit", "Description"], [["KGM", "kilogram"], ["DTN", "100 kilograms"]])
    sheet("Monetary_Units", ["Monetary unit", "Description"], [["EUR", "Euro"]])
    sheet("Measure_Types", ["Measure type", "Description"], [["103", "Third country duty"], ["551", "Synthetic specific measure"]])
    sheet("Duty_Expressions", ["Duty expression", "Description", "Label"], [["01", "Amount in relation to value or quantity", "%"]])
    wb.save(path)
    wb.close()
    if corrupt:
        path.write_bytes(b"not an xlsx workbook")
    return path


def _snapshot(tmp_path: Path) -> tuple[OfficialTaricStore, dict]:
    book = _official_shaped_workbook(tmp_path / "synthetic-test-only.xlsx")
    store = OfficialTaricStore(tmp_path / "store")
    result = store.ingest(book, reference_date="2025-06-01", retrieved_at=datetime(2025, 6, 2, tzinfo=timezone.utc))
    return store, result


def test_snapshot_manifest_hash_status_and_rows(tmp_path):
    store, result = _snapshot(tmp_path)
    active = result["active"]
    assert active["source"] == "EU_TARIC"
    assert active["publisher"] == "European Commission DG TAXUD"
    assert active["reference_date"] == "2025-06-01"
    assert active["retrieved_at"] == "2025-06-02T00:00:00+00:00"
    assert len(active["sha256"]) == 64
    assert len(active["files"][0]["sha256"]) == 64
    assert active["ingestion_version"] == "1.0"
    assert set(active["row_counts"]) == FULL_REQUIRED_TABLES
    assert result["last_good"]["sha256"] == active["sha256"]


def test_directory_and_zip_packages_ingest_offline(tmp_path):
    source_dir = tmp_path / "package"
    source_dir.mkdir()
    _official_shaped_workbook(source_dir / "taric.xlsx")
    store = OfficialTaricStore(tmp_path / "store")
    first = store.ingest(source_dir, reference_date="2025-06-01")
    archive = tmp_path / "snapshot.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.write(source_dir / "taric.xlsx", "taric.xlsx")
    store2 = OfficialTaricStore(tmp_path / "store2")
    second = store2.ingest(archive, reference_date="2025-06-01")
    assert first["active"]["row_counts"] == second["active"]["row_counts"]
    assert first["active"]["sha256"] == second["active"]["sha256"]


def test_missing_required_tables_rejected_without_changing_active(tmp_path):
    store, before = _snapshot(tmp_path)
    sparse = Workbook()
    ws = sparse.active
    ws.title = "Goods_Nomenclature"
    ws.append(["Goods code", "Hierarchy level", "Is leaf"])
    ws.append(["1234567890", 10, 1])
    bad = tmp_path / "bad.xlsx"
    sparse.save(bad)
    with pytest.raises(TaricSnapshotError, match="Measures"):
        store.ingest(bad, reference_date="2025-07-01")
    assert store.status()["active"]["sha256"] == before["active"]["sha256"]


def test_partial_snapshot_is_explicit_and_scoped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Goods_Nomenclature"
    ws.append(["Goods code", "Hierarchy level", "Is leaf"])
    ws.append(["0101", 4, 0])
    ws = wb.create_sheet("Measures")
    ws.append(["Goods code", "Measure type", "Duty expression"])
    ws.append(["0101", "103", "4.5 %"])
    path = tmp_path / "partial.xlsx"
    wb.save(path)
    result = OfficialTaricStore(tmp_path / "store").ingest(path, reference_date="2025-06-01", partial=True)
    assert result["active"]["partial"] is True
    assert "Measure_Exclusions" in result["active"]["missing_tables"]
    assert result["active"]["supported_scope"] == ["nomenclature", "raw measures"]


def test_malformed_workbook_rejected_and_last_good_preserved(tmp_path):
    store, first = _snapshot(tmp_path)
    malformed = tmp_path / "malformed.xlsx"
    malformed.write_bytes(b"bad zip bytes")
    with pytest.raises(TaricSnapshotError, match="Malformed Excel"):
        store.ingest(malformed, reference_date="2025-07-01")
    status = store.status()
    assert status["active"]["sha256"] == first["active"]["sha256"]
    assert status["last_good"]["sha256"] == first["active"]["sha256"]


def test_failed_activation_keeps_active_and_previous_good(tmp_path, monkeypatch):
    store, first = _snapshot(tmp_path)
    second = _official_shaped_workbook(tmp_path / "second.xlsx")
    original = store._sanity_check
    def fail(*args, **kwargs):
        raise TaricSnapshotError("forced validation error")
    monkeypatch.setattr(store, "_sanity_check", fail)
    with pytest.raises(TaricSnapshotError, match="forced"):
        store.ingest(second, reference_date="2025-07-01")
    assert store.status()["active"]["sha256"] == first["active"]["sha256"]
    monkeypatch.setattr(store, "_sanity_check", original)


def test_failed_refresh_marks_last_good_stale(tmp_path):
    store, first = _snapshot(tmp_path)
    store.record_refresh_failure("snapshot_validation_failed", error_code="TaricSnapshotError")
    assert store.status()["state"] == "stale"
    assert store.status()["active"]["sha256"] == first["active"]["sha256"]
    result = store.query("8508700010", "CN", "2025-06-01")
    assert result["stale"] is True
    assert result["source_detail"] == "stale_cache_after_refresh_failure"


def test_query_parent_inheritance_geography_exclusion_and_validity(tmp_path):
    store, _ = _snapshot(tmp_path)
    result = store.query("8508700010", "CN", "2025-06-01")
    assert result["classification"] == {
        "code": "8508700010", "level": "taric10", "is_declarable": True,
        "has_children": False, "required_precision": None, "candidate_subdivisions": [],
    }
    assert len(result["measures"]) == 2
    inherited = next(m for m in result["measures"] if m["measure_type"] == "103")
    specific = next(m for m in result["measures"] if m["measure_type"] == "551")
    assert inherited["inherited_from_parent"] is True
    assert inherited["duty_expression"] == "4.5 %"
    assert inherited["origin_applicability"] == "excluded"
    assert specific["origin_applicability"] == "applicable"
    assert specific["conditions"][0]["condition_code"] == "Y"
    assert specific["duty_expression"] == "60 EUR/tonne"
    assert specific["additional_code_required"] is True
    assert specific["measure_type_description"] == "Synthetic specific measure"
    assert specific["additional_code_details"][0]["description"] == "Synthetic additional code"
    assert specific["legal_basis"][0]["description"] == "Synthetic additional measure reference"
    assert specific["footnotes"][0]["footnote_code"] == "TN001"
    # The earlier, expired leaf measure must not be returned.
    assert all(m["regulation"] != "OLD" for m in result["measures"])


def test_code_precision_uses_official_leaf_data(tmp_path):
    store, _ = _snapshot(tmp_path)
    result = store.query("850870", "CN", "2025-06-01")
    assert result["classification"]["level"] == "hs6"
    assert result["classification"]["is_declarable"] is False
    assert result["classification"]["has_children"] is True
    assert result["classification"]["required_precision"] == "official_declarable_leaf"
    assert "8508700010" in result["classification"]["candidate_subdivisions"]


def test_missing_snapshot_is_source_unavailable(tmp_path):
    result = OfficialTaricStore(tmp_path).query("8508700010", "CN")
    assert result["source"] == "EU_TARIC"
    assert result["source_outcome"] == "source_unavailable"
    assert result["source_detail"] == "snapshot_not_installed"


def test_cli_ingest_and_status(tmp_path):
    book = _official_shaped_workbook(tmp_path / "cli-test.xlsx")
    store_dir = tmp_path / "store"
    ingest = subprocess.run([sys.executable, "-m", "factrail.trade.taric_sync", "ingest", str(book), "--reference-date", "2025-06-01", "--store-dir", str(store_dir)], capture_output=True, text=True)
    assert ingest.returncode == 0, ingest.stderr
    payload = json.loads(ingest.stdout)
    assert payload["active"]["source"] == "EU_TARIC"
    status = subprocess.run([sys.executable, "-m", "factrail.trade.taric_sync", "status", "--store-dir", str(store_dir)], capture_output=True, text=True)
    assert status.returncode == 0
    assert json.loads(status.stdout)["state"] == "ready"


def test_official_store_does_not_claim_network_fetch(tmp_path):
    result = subprocess.run([sys.executable, "-m", "factrail.trade.taric_sync", "fetch", "--store-dir", str(tmp_path)], capture_output=True, text=True)
    assert result.returncode == 2
    assert json.loads(result.stdout)["source_outcome"] == "source_unavailable"


def test_access2markets_compat_module_is_explicitly_secondary():
    from factrail.trade.access2markets_store import Access2MarketsStore
    from factrail.trade.taric_store import TaricStore
    assert TaricStore is Access2MarketsStore
    assert "never an official TARIC" in Access2MarketsStore.__doc__


def test_trade_authority_policies_match_canonical_evidence_fields():
    from factrail.evidence.authority import authority_policy_registry
    policies = {p.policy_id: p for p in authority_policy_registry()}
    rule = policies["trade_import.eu_nomenclature.official_over_heuristic"]
    assert rule.field == "taric_code"
    assert [value.value for value in rule.precedence] == ["official_derived_dataset", "heuristic_classifier"]


def test_old_taric_cache_file_is_migrated_to_access2markets_identity(tmp_path):
    from factrail.trade.access2markets_store import Access2MarketsStore
    db_path = tmp_path / "legacy.db"
    with sqlite3.connect(db_path) as db:
        db.execute("""CREATE TABLE taric_cache (
            hs_code TEXT NOT NULL, origin TEXT NOT NULL, destination TEXT NOT NULL,
            retrieved_at REAL NOT NULL, ttls_until REAL NOT NULL, version TEXT,
            version_date TEXT, measures TEXT NOT NULL, source_note TEXT NOT NULL,
            PRIMARY KEY(hs_code, origin, destination))""")
        db.execute("INSERT INTO taric_cache VALUES (?,?,?,?,?,?,?,?,?)", (
            "850870", "CN", "FR", 4102444800.0, 4102444800.0, "old", "2025-06-01", "[]", "legacy A2M cache"))
    store = Access2MarketsStore(str(db_path))
    with sqlite3.connect(db_path) as db:
        names = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert "access2markets_cache" in names
    assert store.lookup("850870", "CN", "FR")["measures"] == []


def test_evidence_source_identity_and_receipt_metadata(tmp_path, monkeypatch):
    from factrail.evidence.receipts import ReceiptRepository
    from factrail.evidence.trade import assess_import_evidence
    from factrail.trade.models import AssessImportInput
    store, _ = _snapshot(tmp_path)
    monkeypatch.setattr("factrail.trade.service.OfficialTaricStore", lambda: store)
    class EmptySecondary:
        def get_measures(self, *args, **kwargs):
            return {"measures": [], "source_detail": "parser_no_usable_measures", "_source": "live"}
    monkeypatch.setattr("factrail.trade.service.get_default_store", lambda: EmptySecondary())
    parameters = AssessImportInput(product="synthetic test good", origin_country="CN", destination_country="FR",
        quantity=1, goods_value=20, currency="EUR", known_hs_code="8508700010", assessment_date="2025-06-01")
    envelope = assess_import_evidence(parameters)
    official = [source for source in envelope.evidence if source.source_type == "official_customs_source"]
    secondary = [source for source in envelope.evidence if source.source_type == "access2markets_secondary"]
    assert official and secondary
    assert all(source.publisher != "Access2Markets" for source in official)
    assert all(source.authority_class != "official_government_customs_source" for source in secondary)
    measure_fact = next(f for f in envelope.facts if f.field == "customs_measures")
    assert measure_fact.support_level.value == "authoritative"
    assert measure_fact.metadata["snapshot"]["sha256"] == store.status()["active"]["sha256"]
    saved = ReceiptRepository(str(tmp_path / "receipt.sqlite3")).save(envelope)
    loaded = ReceiptRepository(str(tmp_path / "receipt.sqlite3")).get(saved.receipt_id)
    assert loaded is not None
    loaded_fact = next(f for f in loaded.facts if f.field == "customs_measures")
    assert loaded_fact.metadata == measure_fact.metadata
